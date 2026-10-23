import asyncio
from datetime import UTC, datetime
import io
import logging
from aiogram import Bot, F, Router
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
)
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from packages.ieltix_core.adapters.database import Database
from packages.ieltix_core.adapters.repositories import (
    SqlAlchemyAttemptRepository,
    SqlAlchemyProfileRepository,
    SqlAlchemySubscriptionRepository,
    SqlAlchemyUserRepository,
)
from packages.ieltix_core.config import get_settings

logger = logging.getLogger(__name__)

router = Router(name="admin")


class AdminState(StatesGroup):
    waiting_for_broadcast_message = State()
    confirm_broadcast = State()
    waiting_for_new_admin_id = State()
    waiting_for_groq_key = State()
    waiting_for_gemini_key = State()


async def is_admin_user(telegram_id: int, database: Database) -> bool:
    settings = get_settings()
    if telegram_id in settings.admin_ids:
        return True
    async with database.session() as session:
        user = await SqlAlchemyUserRepository(session).get_by_telegram_id(telegram_id)
        return bool(user and user.is_admin)


async def notify_admins(bot: Bot, database: Database, text: str) -> None:
    """Delivers instant real-time notification to all admins."""
    settings = get_settings()
    admin_ids = set(settings.admin_ids)
    try:
        async with database.session() as session:
            db_admin_ids = await SqlAlchemyUserRepository(session).get_all_admin_ids()
            admin_ids.update(db_admin_ids)
    except Exception as exc:
        logger.warning("Could not fetch DB admin ids: %s", exc)

    for admin_id in admin_ids:
        try:
            await bot.send_message(admin_id, text, parse_mode="HTML")
        except Exception as exc:
            logger.warning("Failed to notify admin %s: %s", admin_id, exc)


def admin_dashboard_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="📊 Excel Ro'yxat (.xlsx)", callback_data="admin:export_excel"),
                InlineKeyboardButton(text="📢 Hammaga Xabar (Broadcast)", callback_data="admin:broadcast_start"),
            ],
            [
                InlineKeyboardButton(text="🔑 API Kalitlar (Groq & Gemini)", callback_data="admin:api_keys"),
                InlineKeyboardButton(text="👮‍♂️ Admin Tayinlash", callback_data="admin:add_admin"),
            ],
            [
                InlineKeyboardButton(text="🔄 Yangilash", callback_data="admin:refresh"),
            ],
        ]
    )


async def render_admin_dashboard(database: Database) -> str:
    from packages.ieltix_core.application.key_pool import get_key_pool

    async with database.session() as session:
        user_repo = SqlAlchemyUserRepository(session)
        users = await user_repo.get_all_users()
        total_users = len(users)

        sub_repo = SqlAlchemySubscriptionRepository(session)
        active_subs = 0
        for u in users:
            sub = await sub_repo.get_active(u.id)
            if sub:
                active_subs += 1

    pool = get_key_pool(database)
    keys_info = await pool.get_all_keys_info()
    groq_count = len(keys_info.get("groq", []))
    gemini_count = len(keys_info.get("gemini", []))

    return (
        "👑 <b>MockLife Boshqaruv Paneli (Admin)</b>\n\n"
        f"📊 <b>Umumiy Statistika:</b>\n"
        f"• 👥 <b>Jami foydalanuvchilar:</b> {total_users} ta\n"
        f"• ⭐️ <b>Faol Plus obunachilar:</b> {active_subs} ta\n"
        f"• 🔑 <b>AI API Hovuzi (Pool):</b> ⚡️ {groq_count} ta Groq | ♊️ {gemini_count} ta Gemini\n"

        f"• 👤 <b>Asosiy Admin ID:</b> <code>6956456422</code>\n\n"
        "👇 <i>Quyidagi amallardan birini tanlang:</i>"
    )



@router.message(Command("admin"))
@router.message(Command("panel"))
@router.message(F.text.in_({"👑 Admin Panel", "Admin Panel", "admin", "panel"}))
async def show_admin_panel(message: Message, database: Database) -> None:
    telegram_id = message.from_user.id if message.from_user else 0
    if not await is_admin_user(telegram_id, database):
        await message.answer("⛔️ Kechirasiz, sizda admin huquqi yo'q.")
        return

    text = await render_admin_dashboard(database)
    await message.answer(text, reply_markup=admin_dashboard_keyboard(), parse_mode="HTML")


@router.callback_query(F.data == "admin:refresh")
async def refresh_admin_dashboard(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat berilmagan", show_alert=True)
        return

    text = await render_admin_dashboard(database)
    if callback.message:
        try:
            await callback.message.edit_text(text, reply_markup=admin_dashboard_keyboard(), parse_mode="HTML")
        except Exception:
            pass
    await callback.answer("Yangilandi!")


@router.callback_query(F.data == "admin:export_excel")
async def export_users_excel(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await callback.answer("Excel fayl tayyorlanmoqda...")
    if callback.message:
        await callback.message.bot.send_chat_action(chat_id=callback.message.chat.id, action="upload_document")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Foydalanuvchilar"

    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    headers = [
        "№",
        "Telegram ID",
        "Ism-Familiya",
        "Username",
        "Telefon Raqam",
        "Ro'yxatdan o'tgan sana",
        "Maqsadli Band",
        "Hozirgi Band",
        "Diagnostik Ball",
        "Obuna Holati",
        "Obuna Tugash Sanasi",
        "Admin",
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    now = datetime.now(UTC)
    async with database.session() as session:
        user_repo = SqlAlchemyUserRepository(session)
        profile_repo = SqlAlchemyProfileRepository(session)
        sub_repo = SqlAlchemySubscriptionRepository(session)

        users = await user_repo.get_all_users()
        row_idx = 2
        for idx, u in enumerate(users, 1):
            profile = await profile_repo.get_by_user_id(u.id)
            sub = await sub_repo.get_active(u.id)

            sub_status = "Free"
            sub_expires = "-"
            if sub and sub.valid_until:
                vu = sub.valid_until if sub.valid_until.tzinfo is not None else sub.valid_until.replace(tzinfo=UTC)
                if vu > now:
                    sub_status = "Plus (Faol)"
                    sub_expires = vu.strftime("%d.%m.%Y %H:%M")

            created_str = u.created_at.strftime("%d.%m.%Y %H:%M") if hasattr(u.created_at, "strftime") else "-"
            target_band = getattr(profile, "target_band", "-") if profile else "-"
            current_band = getattr(profile, "current_band", "-") if profile else "-"
            diag_score = getattr(profile, "diagnostic_score", 0) if profile else 0

            row_data = [
                idx,
                u.telegram_id,
                u.full_name,
                f"@{u.username}" if u.username else "-",
                u.phone_number or "-",
                created_str,
                target_band,
                current_band,
                diag_score,
                sub_status,
                sub_expires,
                "Ha" if (u.is_admin or u.telegram_id == 6956456422) else "Yo'q",
            ]
            ws.append(row_data)

            # Apply borders and formatting
            for col_num in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.border = border_thin
                if col_num in (1, 2, 5, 6, 7, 8, 9, 10, 11, 12):
                    c.alignment = center_align
            row_idx += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()
    buf.close()

    filename = f"MockLife_Foydalanuvchilar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    if callback.message:
        await callback.message.answer_document(
            BufferedInputFile(file_bytes, filename=filename),
            caption=f"📊 <b>MockLife Barcha Foydalanuvchilar Ro'yxati</b>\n\nJami: <b>{len(users)}</b> ta foydalanuvchi ma'lumotlari.",
            parse_mode="HTML",
        )


@router.callback_query(F.data == "admin:broadcast_start")
async def start_broadcast(callback: CallbackQuery, state: FSMContext, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await state.set_state(AdminState.waiting_for_broadcast_message)
    await callback.message.answer(
        "📢 <b>Hammaga Xabar Yuborish Rejimi (Broadcast)</b>\n\n"
        "Barcha foydalanuvchilarga yubormoqchi bo'lgan xabaringizni (matn, rasm yoki audio) yozing:\n\n"
        "<i>Bekor qilish uchun /cancel deb yozing.</i>",
        parse_mode="HTML",
    )
    await callback.answer()


@router.message(AdminState.waiting_for_broadcast_message, Command("cancel"))
async def cancel_broadcast(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("❌ Xabar yuborish bekor qilindi.")


@router.message(AdminState.waiting_for_broadcast_message)
async def process_broadcast_message(message: Message, state: FSMContext) -> None:
    await state.update_data(
        broadcast_text=message.text or message.caption or "",
        broadcast_photo_id=message.photo[-1].file_id if message.photo else None,
    )
    await state.set_state(AdminState.confirm_broadcast)

    preview_text = message.text or message.caption or "(Rasm)"
    confirm_kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Ha, yuborilsin!", callback_data="admin:broadcast_confirm"),
                InlineKeyboardButton(text="❌ Bekor qilish", callback_data="admin:broadcast_cancel"),
            ]
        ]
    )
    await message.answer(
        f"⚠️ <b>Xabarni barcha foydalanuvchilarga yuborishni tasdiqlaysizmi?</b>\n\n"
        f"<b>Xabar matni:</b>\n{preview_text[:500]}",
        reply_markup=confirm_kb,
        parse_mode="HTML",
    )


@router.callback_query(AdminState.confirm_broadcast, F.data == "admin:broadcast_cancel")
async def broadcast_canceled(callback: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await callback.message.edit_text("❌ Xabar yuborish bekor qilindi.")
    await callback.answer()


@router.callback_query(AdminState.confirm_broadcast, F.data == "admin:broadcast_confirm")
async def broadcast_confirmed(callback: CallbackQuery, state: FSMContext, database: Database) -> None:
    data = await state.get_data()
    text = data.get("broadcast_text", "")
    photo_id = data.get("broadcast_photo_id")
    await state.clear()

    await callback.message.edit_text("🚀 <b>Xabar yuborilmoqda, kuting...</b>", parse_mode="HTML")
    await callback.answer()

    async with database.session() as session:
        users = await SqlAlchemyUserRepository(session).get_all_users()

    success_count = 0
    fail_count = 0

    for u in users:
        try:
            if photo_id:
                await callback.bot.send_photo(chat_id=u.telegram_id, photo=photo_id, caption=text, parse_mode="HTML")
            else:
                await callback.bot.send_message(chat_id=u.telegram_id, text=text, parse_mode="HTML")
            success_count += 1
            await asyncio.sleep(0.04)  # Anti-flood rate limit
        except Exception:
            fail_count += 1

    await callback.message.answer(
        f"✅ <b>Xabar yuborish yakunlandi!</b>\n\n"
        f"• 📤 <b>Yetkazildi:</b> {success_count} ta foydalanuvchiga\n"
        f"• ⚠️ <b>Yetkazilmadi (bloklaganlar):</b> {fail_count} ta",
        parse_mode="HTML",
    )


@router.callback_query(F.data == "admin:add_admin")
async def add_admin_start(callback: CallbackQuery, state: FSMContext, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await state.set_state(AdminState.waiting_for_new_admin_id)
    await callback.message.answer(
        "👮‍♂️ <b>Yangi Admin Tayinlash:</b>\n\n"
        "Yangi admin qilmoqchi bo'lgan foydalanuvchining <b>Telegram ID</b> raqamini yozing:\n"
        "<i>(Masalan: 6956456422)</i>\n\n"
        "Bekor qilish uchun /cancel deb yozing.",
        parse_mode="HTML",
    )
    await callback.answer()


@router.message(AdminState.waiting_for_new_admin_id, Command("cancel"))
async def cancel_add_admin(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("❌ Bekor qilindi.")


@router.message(AdminState.waiting_for_new_admin_id, F.text)
async def process_new_admin_id(message: Message, state: FSMContext, database: Database) -> None:
    raw_text = message.text.strip()
    if not raw_text.isdigit():
        await message.answer("⚠️ Iltimos, faqat raqamlardan iborat Telegram ID kiriting:")
        return

    target_id = int(raw_text)
    await state.clear()

    async with database.session() as session:
        user_repo = SqlAlchemyUserRepository(session)
        user = await user_repo.get_by_telegram_id(target_id)
        if not user:
            await message.answer(
                f"⚠️ ID <code>{target_id}</code> bo'lgan foydalanuvchi bot bazasida topilmadi. "
                f"U avval botga /start bosgan bo'lishi kerak.",
                parse_mode="HTML",
            )
            return

        await user_repo.set_admin_by_telegram_id(target_id, is_admin=True)

    await message.answer(
        f"✅ <b>Foydalanuvchi muvaffaqiyatli admin qilindi!</b>\n\n"
        f"• 👤 <b>Ism:</b> {user.full_name}\n"
        f"• 🆔 <b>Telegram ID:</b> <code>{target_id}</code>\n"
        f"• 🏷 <b>Username:</b> @{user.username or 'mavjud emas'}",
        parse_mode="HTML",
    )

    try:
        await message.bot.send_message(
            target_id,
            "🎉 <b>Tabriklaymiz!</b> Sizga MockLife botida Admin huquqi berildi.\n/admin buyrug'i orqali boshqaruv panelini ochishingiz mumkin.",
            parse_mode="HTML",
        )
    except Exception:
        pass


def api_keys_dashboard_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="⚡️ + Groq API qo'shish", callback_data="admin:add_groq_key"),
                InlineKeyboardButton(text="♊️ + Gemini API qo'shish", callback_data="admin:add_gemini_key"),
            ],
            [
                InlineKeyboardButton(text="🧪 Kalitlarni Tekshirish", callback_data="admin:test_keys"),
                InlineKeyboardButton(text="🗑 Kalitni O'chirish", callback_data="admin:delete_keys_menu"),
            ],
            [
                InlineKeyboardButton(text="⬅️ Admin Panelga Qaytish", callback_data="admin:back_to_panel"),
            ],
        ]
    )


async def render_api_keys_text(database: Database) -> str:
    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    keys_info = await pool.get_all_keys_info()
    groq_list = keys_info.get("groq", [])
    gemini_list = keys_info.get("gemini", [])

    text_parts = [
        "🔑 <b>API Kalitlar Hovuzi (Multi-Key Pool)</b>\n",
        "<i>Bir nechta foydalanuvchi bir vaqtda foydalanganda limit (HTTP 429) tugab qolmasligi uchun barcha so'rovlar ushbu kalitlar o'rtasida avtomatik navbatma-navbat (load-balancing & failover) taqsimlanadi.</i>\n",
        f"⚡️ <b>Groq API kalitlari ({len(groq_list)} ta):</b>",
    ]

    if groq_list:
        for idx, k in enumerate(groq_list, 1):
            status = "⏳ Kutishda (429)" if k.get("is_cooldown") else ("🟢 Faol" if k.get("is_active") else "🔴 O'chiq")
            text_parts.append(
                f"{idx}. <code>{k['masked']}</code> — {k.get('label') or 'Groq Key'}\n"
                f"   Holati: {status} | Jami: {k.get('total_calls', 0)} ta | Xato: {k.get('failed_calls', 0)} ta"
            )
    else:
        text_parts.append("<i>Hozircha Groq kalitlari qo'shilmagan.</i>")

    text_parts.append(f"\n♊️ <b>Gemini API kalitlari ({len(gemini_list)} ta):</b>")
    if gemini_list:
        for idx, k in enumerate(gemini_list, 1):
            status = "⏳ Kutishda (429)" if k.get("is_cooldown") else ("🟢 Faol" if k.get("is_active") else "🔴 O'chiq")
            text_parts.append(
                f"{idx}. <code>{k['masked']}</code> — {k.get('label') or 'Gemini Key'}\n"
                f"   Holati: {status} | Jami: {k.get('total_calls', 0)} ta | Xato: {k.get('failed_calls', 0)} ta"
            )
    else:
        text_parts.append("<i>Hozircha Gemini kalitlari qo'shilmagan.</i>")

    text_parts.append("\n👇 <i>Quyidagi amallardan birini tanlang:</i>")
    return "\n".join(text_parts)


@router.callback_query(F.data == "admin:api_keys")
async def show_api_keys(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    text = await render_api_keys_text(database)
    if callback.message:
        try:
            await callback.message.edit_text(text, reply_markup=api_keys_dashboard_keyboard(), parse_mode="HTML")
        except Exception:
            await callback.message.answer(text, reply_markup=api_keys_dashboard_keyboard(), parse_mode="HTML")
    await callback.answer()


@router.callback_query(F.data == "admin:back_to_panel")
async def back_to_admin_panel(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    text = await render_admin_dashboard(database)
    if callback.message:
        try:
            await callback.message.edit_text(text, reply_markup=admin_dashboard_keyboard(), parse_mode="HTML")
        except Exception:
            await callback.message.answer(text, reply_markup=admin_dashboard_keyboard(), parse_mode="HTML")
    await callback.answer()


@router.callback_query(F.data == "admin:add_groq_key")
async def add_groq_key_start(callback: CallbackQuery, state: FSMContext, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await state.set_state(AdminState.waiting_for_groq_key)
    await callback.message.answer(
        "⚡️ <b>Yangi Groq API Kalitini Qo'shish:</b>\n\n"
        "Groq Cloud konsolidan (https://console.groq.com/keys) nusxalangan API kalitni yuboring:\n"
        "<i>(Odatda <code>gsk_...</code> bilan boshlanadi)</i>\n\n"
        "Bekor qilish uchun /cancel deb yozing.",
        parse_mode="HTML",
    )
    await callback.answer()


@router.message(AdminState.waiting_for_groq_key, Command("cancel"))
async def cancel_add_groq_key(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("❌ Groq API kalit qo'shish bekor qilindi.")


@router.message(AdminState.waiting_for_groq_key, F.text)
async def process_groq_key(message: Message, state: FSMContext, database: Database) -> None:
    key_input = message.text.strip()
    await state.clear()

    await message.answer("🔄 <i>Groq API kaliti tekshirilmoqda va ulanmoqda...</i>", parse_mode="HTML")

    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    success, msg = await pool.add_key(provider="groq", api_key=key_input, label=f"Groq Qo'shimcha ({datetime.now().strftime('%d.%m')})")

    if not success:
        await message.answer(
            f"❌ <b>Groq kaliti qo'shilmadi:</b>\n{msg}\n\n"
            "Iltimos, kalit to'g'riligini tekshirib qaytadan urinib ko'ring.",
            parse_mode="HTML",
        )
        return

    dash_text = await render_api_keys_text(database)
    await message.answer(
        f"✅ <b>Yangi Groq API kaliti muvaffaqiyatli qo'shildi va faollashtirildi!</b>\n\n"
        f"• Natija: {msg}\n"
        f"• Endi barcha Speaking va STT so'rovlari ushbu kalitlar bo'yicha avtomatik taqsimlanadi.\n\n"
        f"{dash_text}",
        reply_markup=api_keys_dashboard_keyboard(),
        parse_mode="HTML",
    )


@router.callback_query(F.data == "admin:add_gemini_key")
async def add_gemini_key_start(callback: CallbackQuery, state: FSMContext, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await state.set_state(AdminState.waiting_for_gemini_key)
    await callback.message.answer(
        "♊️ <b>Yangi Google Gemini API Kalitini Qo'shish:</b>\n\n"
        "Google AI Studio (https://aistudio.google.com/app/apikey) dan olingan Gemini API kalitni yuboring:\n"
        "<i>(Odatda <code>AIzaSy...</code> bilan boshlanadi)</i>\n\n"
        "Bekor qilish uchun /cancel deb yozing.",
        parse_mode="HTML",
    )
    await callback.answer()


@router.message(AdminState.waiting_for_gemini_key, Command("cancel"))
async def cancel_add_gemini_key(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("❌ Gemini API kalit qo'shish bekor qilindi.")


@router.message(AdminState.waiting_for_gemini_key, F.text)
async def process_gemini_key(message: Message, state: FSMContext, database: Database) -> None:
    key_input = message.text.strip()
    await state.clear()

    await message.answer("🔄 <i>Gemini API kaliti tekshirilmoqda va ulanmoqda...</i>", parse_mode="HTML")

    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    success, msg = await pool.add_key(provider="gemini", api_key=key_input, label=f"Gemini Qo'shimcha ({datetime.now().strftime('%d.%m')})")

    if not success:
        await message.answer(
            f"❌ <b>Gemini kaliti qo'shilmadi:</b>\n{msg}\n\n"
            "Iltimos, kalit to'g'riligini tekshirib qaytadan urinib ko'ring.",
            parse_mode="HTML",
        )
        return

    dash_text = await render_api_keys_text(database)
    await message.answer(
        f"✅ <b>Yangi Gemini API kaliti muvaffaqiyatli qo'shildi va faollashtirildi!</b>\n\n"
        f"• Natija: {msg}\n\n"
        f"{dash_text}",
        reply_markup=api_keys_dashboard_keyboard(),
        parse_mode="HTML",
    )


@router.callback_query(F.data == "admin:test_keys")
async def test_all_api_keys(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    await callback.answer("Barcha kalitlar tekshirilmoqda, kuting...")
    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    keys_info = await pool.get_all_keys_info()

    results = ["🧪 <b>API Kalitlarni Jonli Tekshirish Natijasi:</b>\n"]

    groq_list = keys_info.get("groq", [])
    results.append("⚡️ <b>Groq Kalitlari:</b>")
    if groq_list:
        for k in groq_list:
            is_ok, msg = await pool.test_groq_key(k["api_key"])
            status_icon = "🟢 200 OK (Ishlayapti)" if is_ok else f"🔴 Xato ({msg})"
            results.append(f"• <code>{k['masked']}</code>: {status_icon}")
    else:
        results.append("<i>Groq kalitlari yo'q</i>")

    gemini_list = keys_info.get("gemini", [])
    results.append("\n♊️ <b>Gemini Kalitlari:</b>")
    if gemini_list:
        for k in gemini_list:
            is_ok, msg = await pool.test_gemini_key(k["api_key"])
            status_icon = "🟢 200 OK (Ishlayapti)" if is_ok else f"🔴 Xato ({msg})"
            results.append(f"• <code>{k['masked']}</code>: {status_icon}")
    else:
        results.append("<i>Gemini kalitlari yo'q</i>")

    results.append("\n<i>Kalitlar avtomatik navbatma-navbat ishlamoqda.</i>")
    test_text = "\n".join(results)

    test_kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🔄 Qayta Tekshirish", callback_data="admin:test_keys")],
            [InlineKeyboardButton(text="⬅️ Kalitlar Menyusiga Qaytish", callback_data="admin:api_keys")],
        ]
    )

    if callback.message:
        try:
            await callback.message.edit_text(test_text, reply_markup=test_kb, parse_mode="HTML")
        except Exception:
            await callback.message.answer(test_text, reply_markup=test_kb, parse_mode="HTML")


@router.callback_query(F.data == "admin:delete_keys_menu")
async def delete_keys_menu(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    keys_info = await pool.get_all_keys_info()

    buttons = []
    for k in keys_info.get("groq", []):
        buttons.append([
            InlineKeyboardButton(
                text=f"🗑 Groq: {k['masked']} ({k.get('label', '')[:15]})",
                callback_data=f"admin:delete_key:{k['id']}",
            )
        ])
    for k in keys_info.get("gemini", []):
        buttons.append([
            InlineKeyboardButton(
                text=f"🗑 Gemini: {k['masked']} ({k.get('label', '')[:15]})",
                callback_data=f"admin:delete_key:{k['id']}",
            )
        ])

    buttons.append([InlineKeyboardButton(text="⬅️ Bekor qilish", callback_data="admin:api_keys")])

    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    if callback.message:
        try:
            await callback.message.edit_text(
                "🗑 <b>O'chirmoqchi bo'lgan API kalitni tanlang:</b>\n\n"
                "<i>Eslatma: O'chirilgan kalit darhol so'rovlar hovuzidan chiqariladi.</i>",
                reply_markup=kb,
                parse_mode="HTML",
            )
        except Exception:
            pass
    await callback.answer()


@router.callback_query(F.data.startswith("admin:delete_key:"))
async def process_delete_key(callback: CallbackQuery, database: Database) -> None:
    if not await is_admin_user(callback.from_user.id, database):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return

    raw_id = callback.data.split(":")[-1]
    if not raw_id.isdigit():
        await callback.answer("Xatolik", show_alert=True)
        return

    key_id = int(raw_id)
    from packages.ieltix_core.application.key_pool import get_key_pool

    pool = get_key_pool(database)
    deleted = await pool.delete_key_by_id(key_id)

    if deleted:
        await callback.answer("Kalit muvaffaqiyatli o'chirildi!", show_alert=True)
    else:
        await callback.answer("Kalit topilmadi yoki o'chirilmadi.", show_alert=True)

    text = await render_api_keys_text(database)
    if callback.message:
        try:
            await callback.message.edit_text(text, reply_markup=api_keys_dashboard_keyboard(), parse_mode="HTML")
        except Exception:
            pass

